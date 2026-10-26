"""Package the reviewed Vietnamese PDF as a layout-preserving presentation.

Slide artwork is a raster snapshot, not editable PowerPoint text. Edit the
HTML/CSS source and rebuild instead. Vietnamese speaker notes remain editable.
No application, credentials or external service is accessed.
"""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import fitz
from bs4 import BeautifulSoup
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.util import Inches


def build() -> None:
    root = Path(__file__).resolve().parents[1]
    output = root / "docs/submission"
    stem = "ExaMate-Technologies-VI-20261026"
    pdf_path = output / f"{stem}.pdf"
    source = BeautifulSoup(
        (root / "docs/presentation/technology.html").read_text(encoding="utf-8"),
        "html.parser",
    )
    sections = source.select("main > section")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    properties = presentation.core_properties
    properties.title = "ExaMate — Công nghệ và quyết định thiết kế"
    properties.subject = "Bản tiếng Việt; artwork cố định, lời nói có thể chỉnh sửa"
    properties.author = "Nhóm ExaMate — Tài và Thắng"
    properties.identifier = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    thumbnails = []
    with fitz.open(pdf_path) as document:
        if len(document) != len(sections) or len(document) != 15:
            raise ValueError("PDF and HTML must describe the same 15 slides")
        for index, (page, section) in enumerate(zip(document, sections)):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            data = io.BytesIO()
            image.save(data, format="JPEG", quality=95, subsampling=0)
            data.seek(0)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            artwork = slide.shapes.add_picture(
                data, 0, 0, presentation.slide_width, presentation.slide_height
            )
            title = section.select_one("h1,h2").get_text(" ", strip=True)
            artwork._element.nvPicPr.cNvPr.set("descr", f"{index + 1:02}. {title}")
            slide.notes_slide.notes_text_frame.text = (
                section.select_one(".speaker").get_text("\n", strip=True)
                + "\n\nBố cục cố định từ PDF; chỉnh nội dung trong technology.html/CSS rồi dựng lại."
            )
            tile = Image.new("RGB", (504, 304), "#e5e5e0")
            thumbnail = ImageOps.contain(image, (480, 270))
            tile.paste(thumbnail, ((504 - thumbnail.width) // 2, 8))
            ImageDraw.Draw(tile).text((12, 282), f"{index + 1:02} / 15", fill="#191b1e")
            thumbnails.append(tile)
    presentation.save(output / f"{stem}.pptx")
    preview = Image.new("RGB", (1512, 1520), "#dededb")
    for index, image in enumerate(thumbnails):
        preview.paste(image, ((index % 3) * 504, (index // 3) * 304))
    preview.save(output / f"{stem}-preview.jpg", quality=88)
    manifest_path = output / "technology-deck-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest.update({
        "pptx": f"{stem}.pptx",
        "pptxSha256": hashlib.sha256((output / f"{stem}.pptx").read_bytes()).hexdigest(),
        "pptxArtwork": "raster-snapshot-from-pdf; editable speaker notes",
        "preview": f"{stem}-preview.jpg",
        "previewSha256": hashlib.sha256((output / f"{stem}-preview.jpg").read_bytes()).hexdigest(),
    })
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("Built layout-preserving PPTX: 15 slides with Vietnamese speaker notes; HTML is editable source.")


if __name__ == "__main__":
    build()
