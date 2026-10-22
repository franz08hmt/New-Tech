"""Builds the submission documents from their editable sources.

  python scripts/build-submission-docs.py <output-dir>

Sources (tracked): docs/REPORT.md, docs/architecture.svg, docs/SLIDES.md.
Outputs: ExaMate-report-DRAFT.pdf, ExaMate-slides-DRAFT.pptx, ExaMate-slides-DRAFT.pdf.

Needs PyMuPDF (fitz), python-pptx and markdown; none is a dependency of the
application. The PDFs are renderings of the Markdown: edit the Markdown, not
the PDF. Run it again after every change; the page count is printed.
"""
import re
import sys
from pathlib import Path

import fitz
import markdown
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.4; color: #1f2a21; }
h1 { font-size: 19pt; color: #2f4a35; margin: 0 0 6pt 0; }
h2 { font-size: 13.5pt; color: #2f4a35; margin: 12pt 0 4pt 0; border-bottom: 1px solid #9db39f; }
h3 { font-size: 11pt; color: #3b4a3d; margin: 9pt 0 3pt 0; }
p { margin: 3pt 0 4pt 0; }
li { margin: 1pt 0; }
code { font-family: monospace; font-size: 8.3pt; background-color: #eef3ec; }
pre { font-family: monospace; font-size: 8pt; background-color: #eef3ec; padding: 4pt; }
blockquote { margin: 4pt 0 6pt 0; padding: 3pt 8pt; background-color: #fbeee4; border-left: 3px solid #b4531c; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 7pt 0; }
th { font-size: 8.6pt; background-color: #dfe8dc; border: 1px solid #8ea592; padding: 2.5pt; text-align: left; }
td { font-size: 8.6pt; border: 1px solid #b9c7bb; padding: 2.5pt; vertical-align: top; }
img { width: 100%; }
"""


def render_architecture(out_dir: Path) -> Path:
    doc = fitz.open(str(DOCS / "architecture.svg"))
    target = out_dir / "architecture.png"
    doc[0].get_pixmap(dpi=150).save(str(target))
    return target


def markdown_to_pdf(source: Path, target: Path, out_dir: Path) -> int:
    html = markdown.markdown(
        source.read_text(encoding="utf-8"),
        extensions=["tables", "fenced_code", "sane_lists"],
    )
    archive = fitz.Archive(str(out_dir))
    story = fitz.Story(html=html, user_css=CSS, archive=archive)
    raw = target.with_suffix(".raw.pdf")
    writer = fitz.DocumentWriter(str(raw))
    page = fitz.paper_rect("a4")
    where = page + (46, 46, -46, -52)
    more = 1
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    doc = fitz.open(str(raw))
    for number, pg in enumerate(doc, start=1):
        pg.insert_text(
            (46, pg.rect.height - 28),
            f"ExaMate AI - bao cao DRAFT - trang {number}/{doc.page_count}",
            fontsize=7.5,
            color=(0.4, 0.4, 0.4),
        )
    count = doc.page_count
    doc.subset_fonts()
    doc.save(str(target), garbage=4, deflate=True)
    doc.close()
    del writer, story
    try:
        raw.unlink()
    except PermissionError:
        pass  # Windows may keep the intermediate file open until exit
    return count


# ---- slides ---------------------------------------------------------------

def parse_slides(path: Path):
    """`## Title` starts a slide; `- ` lines are bullets; `> ` lines are notes;
    `![](file)` is an image; everything else is ignored."""
    slides, current = [], None
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("## "):
            current = {"title": line[3:].strip(), "bullets": [], "notes": [], "image": None}
            slides.append(current)
        elif current is None:
            continue
        elif line.startswith("- "):
            current["bullets"].append(line[2:].strip())
        elif line.startswith("  - "):
            current["bullets"].append("   - " + line[4:].strip())
        elif line.startswith("> "):
            current["notes"].append(line[2:].strip())
        elif line.startswith("!["):
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            if m:
                current["image"] = m.group(1)
    return slides


def plain(text: str) -> str:
    return re.sub(r"[`*]", "", text)


def build_pptx(slides, target: Path, out_dir: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    layout = prs.slide_layouts[6]
    green, ink = RGBColor(0x2F, 0x4A, 0x35), RGBColor(0x1F, 0x2A, 0x21)
    for slide in slides:
        s = prs.slides.add_slide(layout)
        title = s.shapes.add_textbox(Emu(450000), Emu(250000), Emu(11300000), Emu(800000))
        tf = title.text_frame
        tf.word_wrap = True
        tf.text = plain(slide["title"])
        tf.paragraphs[0].font.size = Pt(30)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = green
        has_image = bool(slide["image"])
        width = Emu(5600000) if has_image else Emu(11300000)
        body = s.shapes.add_textbox(Emu(450000), Emu(1150000), width, Emu(5300000))
        bf = body.text_frame
        bf.word_wrap = True
        size = 20 if len(slide["bullets"]) <= 5 else 17 if len(slide["bullets"]) <= 7 else 15
        for i, bullet in enumerate(slide["bullets"]):
            p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            sub = bullet.startswith("   - ")
            p.text = ("– " if sub else "• ") + plain(bullet.lstrip(" -"))
            p.level = 1 if sub else 0
            p.font.size = Pt(size - 3 if sub else size)
            p.font.color.rgb = ink
            p.space_after = Pt(6)
        if has_image:
            png = out_dir / Path(slide["image"]).name
            if slide["bullets"]:
                s.shapes.add_picture(str(png), Emu(6200000), Emu(1250000), width=Emu(5600000))
            else:  # image-only slide: as large as the page allows
                s.shapes.add_picture(str(png), Emu(900000), Emu(1100000), height=Emu(5500000))
        if slide["notes"]:
            s.notes_slide.notes_text_frame.text = "\n".join(plain(n) for n in slide["notes"])
    prs.save(str(target))
    return len(slides)


def find_font(candidates):
    for path in candidates:
        if Path(path).exists():
            return path
    return None


FONT_REGULAR = find_font([
    "C:/Windows/Fonts/arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
])
FONT_BOLD = find_font([
    "C:/Windows/Fonts/arialbd.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
])


def build_slides_pdf(slides, target: Path, out_dir: Path):
    """The PDF needs a TrueType font that has Vietnamese letters; the built-in
    Helvetica does not, so it is only a fallback."""
    doc = fitz.open()
    width, height = 960, 540
    for number, slide in enumerate(slides, start=1):
        page = doc.new_page(width=width, height=height)
        regular, bold = "helv", "hebo"
        if FONT_REGULAR and FONT_BOLD:
            page.insert_font(fontname="vi", fontfile=FONT_REGULAR)
            page.insert_font(fontname="vib", fontfile=FONT_BOLD)
            regular, bold = "vi", "vib"
        page.insert_textbox(fitz.Rect(34, 20, width - 34, 80), plain(slide["title"]),
                            fontsize=26, fontname=bold, color=(0.18, 0.29, 0.21))
        size = 22 if len(slide["bullets"]) <= 5 else 18 if len(slide["bullets"]) <= 7 else 15
        text_w = 450 if slide["image"] else width - 68
        y = 92
        for bullet in slide["bullets"]:
            sub = bullet.startswith("   - ")
            label = ("- " if sub else "* ") + plain(bullet.lstrip(" -"))
            rect = fitz.Rect(34 + (18 if sub else 0), y, 34 + text_w, y + 120)
            rest = page.insert_textbox(rect, label, fontsize=size - (2 if sub else 0),
                                       fontname=regular, color=(0.12, 0.16, 0.13))
            lines = max(1, round(len(label) * (size * 0.5) / (text_w - (18 if sub else 0))) + 1)
            y += lines * (size + 3) + 5
        if slide["image"]:
            png = out_dir / Path(slide["image"]).name
            if slide["bullets"]:
                page.insert_image(fitz.Rect(500, 96, width - 24, 96 + 300), filename=str(png))
            else:
                page.insert_image(fitz.Rect(70, 78, width - 70, height - 30), filename=str(png))
        page.insert_text((34, height - 14), f"{number}/{len(slides)}", fontsize=8, color=(0.5, 0.5, 0.5))
    doc.subset_fonts()
    doc.save(str(target), garbage=4, deflate=True)


def main():
    out_dir = Path(sys.argv[1]).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    render_architecture(out_dir)
    pages = markdown_to_pdf(DOCS / "REPORT.md", out_dir / "ExaMate-report-DRAFT.pdf", out_dir)
    print(f"report pages: {pages}")
    slides_src = DOCS / "SLIDES.md"
    if slides_src.exists():
        slides = parse_slides(slides_src)
        print(f"slides: {build_pptx(slides, out_dir / 'ExaMate-slides-DRAFT.pptx', out_dir)}")
        build_slides_pdf(slides, out_dir / "ExaMate-slides-DRAFT.pdf", out_dir)


if __name__ == "__main__":
    main()
