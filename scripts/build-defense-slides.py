"""Build ExaMate's designed defense deck from docs/SLIDES.md, without network.

Uses the existing outline parser and installed authoring libraries. Text, diagrams
and evaluation bars are editable PowerPoint shapes. PDF is exported by installed
PowerPoint on Windows so it shares the presentation's actual rendering.
No application credentials, database, provider or execution ledger are accessed.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import subprocess
from pathlib import Path

import fitz
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs/SLIDES.md"
ASSETS = ROOT / "docs/report-assets"
EVIDENCE = ROOT / "docs/submission/report-evidence.json"
STEM = "ExaMate-Presentation-20261024"
W, H = 960, 540
COL = dict(bg="F5F3EC", white="FFFFFF", ink="18372F", muted="586A62",
           line="D7DED5", green="335D49", pale="E6EDDF", lime="D2E58B",
           orange="B85B34", sand="F0DFD0", dark="142D29", gray="E2E4DF")
FONTS = {
    ("Arial", False): "arial.ttf", ("Arial", True): "arialbd.ttf",
    ("Georgia", False): "georgia.ttf", ("Georgia", True): "georgiab.ttf",
}
METRICS = {key: fitz.Font(fontfile=str(Path("C:/Windows/Fonts") / value))
           for key, value in FONTS.items()}


def load_outline():
    spec = importlib.util.spec_from_file_location("submission_docs", ROOT / "scripts/build-submission-docs.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    slides = module.parse_slides(SOURCE)
    if len(slides) != 14:
        raise ValueError("The designed deck requires 12 presentation + 2 backup slides.")
    expected = [6, 4, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6]
    if [len(s["bullets"]) for s in slides] != expected or any(not s["notes"] for s in slides):
        raise ValueError("Outline structure or speaker notes changed; review the layouts.")
    return slides


def color(name):
    return RGBColor.from_string(COL.get(name, name))


def wrap(text, width, size, family="Arial", bold=False):
    metric = METRICS[(family, bold)]
    result = []
    for line in text.split("\n"):
        current = ""
        for word in line.split():
            candidate = (current + " " + word).strip()
            if current and metric.text_length(candidate, fontsize=size) > width:
                result.append(current)
                current = word
            else:
                current = candidate
            if metric.text_length(current, fontsize=size) > width:
                raise ValueError(f"Unbreakable text exceeds layout: {current}")
        result.append(current)
    return result


class Deck:
    def __init__(self, outline):
        self.outline = outline
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(W), Pt(H)
        self.prs.core_properties.title = "ExaMate — AI Academic Workspace"
        self.prs.core_properties.subject = "A: RAG · C: Structured Study Practice"
        self.prs.core_properties.author = "ExaMate team — Tài and Thắng"
        self.prs.core_properties.keywords = "DRAFT; retained evidence; defense; 24 October 2026"
        for rel in self.prs.slide_master.part.rels.values():
            if rel.reltype.endswith('/theme'):
                theme = etree.fromstring(rel.target_part.blob)
                for node in theme.xpath('//*[local-name()="hlink" or local-name()="folHlink"]'):
                    for child in list(node):
                        node.remove(child)
                    fill = OxmlElement('a:srgbClr')
                    fill.set('val', COL['dark'])
                    node.append(fill)
                rel.target_part._blob = etree.tostring(theme)
        self.geometry, self.used_images = [], set()

    def slide(self, index, section, speaker, dark=False):
        self.index, self.dark = index, dark
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.s.background.fill.solid()
        self.s.background.fill.fore_color.rgb = color("dark" if dark else "bg")
        self.text(38, 23, 680, 17, section.upper(), 10, "lime" if dark else "muted", True)
        self.text(758, 23, 164, 17, speaker.upper(), 10, "lime" if dark else "muted", True, align="right")
        self.line(38, 497, 922, 497, "green" if dark else "line")
        self.text(38, 509, 820, 15, "EXAMATE  /  DRAFT  /  EVIDENCE: 23 OCT 2026", 9, "pale" if dark else "muted")
        self.text(862, 507, 60, 17, f"{index + 1:02} / 14", 10, "pale" if dark else "muted", align="right")
        self.s.notes_slide.notes_text_frame.text = "\n\n".join(self.outline[index]["notes"])
        return self.outline[index]["bullets"]

    def title(self, text=None, width=880, size=34):
        text = text or self.outline[self.index]["title"]
        self.text(38, 57, width, 83, text, size, "white" if self.dark else "ink", family="Georgia")

    def rect(self, x, y, w, h, fill="white", stroke=None, radius=False):
        shape = self.s.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE if radius else SHAPE.RECTANGLE,
                                        Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
        if stroke:
            shape.line.color.rgb = color(stroke)
            shape.line.width = Pt(.8)
        else:
            shape.line.fill.background()
        if radius:
            shape.adjustments[0] = .08
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        self.geometry.append(dict(slide=self.index + 1, kind="shape", x=x, y=y, w=w, h=h))
        return shape

    def line(self, x1, y1, x2, y2, stroke="line", arrow=False):
        shape = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(1.2)
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        if arrow:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            shape._element.spPr.get_or_add_ln().append(end)

    def text(self, x, y, w, h, text, size=18, fill="ink", bold=False,
             family="Arial", align="left", link=None):
        lines = wrap(text, w - 3, size, family, bold)
        needed = len(lines) * size * 1.16 + 3
        if needed > h:
            raise ValueError(f"Slide {self.index+1}: text needs {needed:.1f}pt, box is {h}pt: {text}")
        shape = self.s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        frame = shape.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = False
        frame.vertical_anchor = MSO_ANCHOR.TOP
        p = frame.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = Pt(size * 1.16)
        p.space_before = p.space_after = Pt(0)
        for i, line in enumerate(lines):
            if i:
                p.add_line_break()
            run = p.add_run()
            run.text = line
            run.font.name, run.font.size = family, Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color(fill)
            if link:
                run.hyperlink.address = link
                run.font.underline = False
        self.geometry.append(dict(slide=self.index+1, kind="text", x=x, y=y, w=w, h=h,
                                  size=size, text=text, lines=len(lines)))
        return shape

    def image(self, path, x, y, w, h):
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        with Image.open(path) as im:
            ratio = min(w/im.width, h/im.height)
            iw, ih = im.width*ratio, im.height*ratio
        self.s.shapes.add_picture(str(path), Pt(x+(w-iw)/2), Pt(y+(h-ih)/2), width=Pt(iw), height=Pt(ih))
        self.used_images.add(path)
        self.geometry.append(dict(slide=self.index+1, kind="image", x=x, y=y, w=w, h=h))

    def pill(self, x, y, w, label, dark=False):
        self.rect(x, y, w, 30, "green" if dark else "pale", radius=True)
        self.text(x+12, y+7, w-24, 20, label, 11, "white" if dark else "green", True)

    def pair(self, text):
        return [p.strip() for p in text.split("|", 1)]

    def step(self, x, y, w, text, n, h=65, dark=False):
        label, body = self.pair(text)
        self.rect(x, y, w, h, "green" if dark else "white", radius=True)
        self.text(x+14, y+12, 25, 28, f"{n:02}", 17, "lime" if dark else "orange", True)
        self.text(x+50, y+10, w-64, 26, label, 17, "white" if dark else "ink", True)
        self.text(x+50, y+31, w-65, h-31, body, 13, "pale" if dark else "muted")


def compose(deck, evidence):
    d = deck
    b = d.slide(0, "New Technologies / project defense", "Tài", dark=True)
    d.pill(40, 110, 216, b[0], True)
    d.text(38, 168, 460, 90, "ExaMate", 69, "white", family="Georgia")
    d.text(42, 270, 385, 76, b[1], 25, "pale")
    d.pill(42, 370, 240, b[4], True)
    d.pill(42, 411, 256, b[5], True)
    d.rect(481, 140, 441, 273, "white", radius=True)
    d.image(ASSETS/"dashboard.jpg", 491, 150, 421, 253)
    d.text(490, 424, 420, 36, b[2] + "  ·  " + b[3], 12, "pale")

    b = d.slide(1, "Problem / product", "Tài")
    d.title()
    for i in range(3):
        d.step(38, 155+i*94, 294, b[i], i+1, 86)
    d.rect(353, 155, 569, 284, "white", radius=True)
    d.image(ASSETS/"dashboard.jpg", 363, 164, 549, 263)
    d.text(359, 444, 557, 34, b[3] + "  Real-service capture, retained.", 12, "muted")

    b = d.slide(2, "System design", "Thắng")
    d.title()
    d.rect(38, 167, 218, 140, "white", "line", True)
    d.text(54, 182, 186, 31, d.pair(b[0])[0], 23, bold=True)
    d.text(54, 226, 186, 65, d.pair(b[0])[1], 17, "muted")
    d.line(256, 234, 308, 234, "orange", True)
    d.rect(309, 151, 299, 274, "ink", radius=True)
    d.text(330, 169, 258, 37, "NestJS API", 28, "white", True)
    d.text(330, 218, 250, 164, "Session + membership\nBusiness services\nRAG / Study Coach\nAdmission + telemetry", 19, "pale")
    d.text(330, 381, 250, 27, "Local Whisper worker", 16, "lime", True)
    for i in range(3):
        label, text = d.pair(b[i+2])
        y = 146+i*99
        d.line(608, y+40, 664, y+40, "green", True)
        d.rect(665, y, 257, 88, "white", "line", True)
        d.text(680, y+10, 227, 26, label, 17, bold=True)
        d.text(680, y+38, 227, 43, text, 14, "muted")
    d.text(38, 454, 880, 32, b[5], 16, "green", True)

    b = d.slide(3, "Capability A / RAG", "Tài")
    d.title(size=32)
    for i in range(4):
        text = b[i].split("·", 1)[1].strip()
        d.step(38, 148+i*71, 394, text, i+1, 65)
    d.rect(451, 148, 471, 290, "white", radius=True)
    d.image(ASSETS/"rag-citation-page3.jpg", 459, 154, 455, 276)
    d.text(453, 442, 464, 39, "Historical browser fixture: real Nest/PG17; synthetic Auth, Storage and provider. PDF page 3 shown.", 10, "muted")
    d.text(38, 443, 391, 44, "No evidence → no answer generation.\nQuery embedding may already have run.", 12, "green", True)

    b = d.slide(4, "Demonstration / switch to the web app", "Tài → Thắng → Tài", True)
    d.title(size=35)
    for i in range(4):
        text = b[i].split("·", 1)[1].strip()
        d.step(38+(i%2)*451, 149+(i//2)*99, 433, text, i+1, 88, True)
    d.text(40, 355, 878, 68, '“' + b[4] + '”', 23, "white", family="Georgia")
    d.rect(38, 443, 884, 39, "lime", radius=True)
    d.text(54, 452, 850, 28, b[5] + "   ↗   OPEN THE LOCAL DEMO", 17, "dark", True, link=b[5])

    b = d.slide(5, "Capability C / structured output", "Tài")
    d.title(size=32)
    for i in range(4):
        d.step(38, 148+i*72, 438, b[i], i+1, 66)
    d.rect(495, 148, 427, 295, "white", radius=True)
    d.image(ASSETS/"quiz-result.png", 507, 158, 403, 274)
    d.text(499, 447, 416, 36, "Historical fixture: real server grading, synthetic provider. Self-rating is not an AI grade.", 10, "muted")
    d.text(38, 447, 440, 35, b[4], 16, "green", True)

    b = d.slide(6, "Streaming / speech input", "Thắng")
    d.title(size=31)
    for i, (x, label) in enumerate([(38, "SSE  /  LIVE RESPONSE"), (489, "STT  /  LOCAL WHISPER")]):
        d.rect(x, 150, 433, 265, "white", "line", True)
        d.text(x+20, 169, 391, 24, label, 12, "orange", True)
        if i == 0:
            d.text(x+20, 211, 389, 99, d.pair(b[0])[1], 23, family="Georgia")
            d.line(x+20, 326, x+411, 326)
            d.text(x+20, 343, 389, 60, d.pair(b[1])[1], 16, "muted")
        else:
            d.text(x+20, 211, 389, 99, "Record → transcribe\nEdit → insert → send", 24, family="Georgia")
            d.line(x+20, 326, x+411, 326)
            d.text(x+20, 342, 389, 61, "60 s audio / 120 s inference\n150 s request / 180 s outer wait", 16, "muted")
    d.text(38, 437, 877, 43, b[5] + " " + b[4], 15, "green", True)

    b = d.slide(7, "Access / validation / recovery", "Tài + Thắng")
    d.title(size=31)
    for i in range(4):
        x, y = 38+(i%2)*451, 147+(i//2)*133
        label, body = d.pair(b[i])
        d.rect(x, y, 433, 120, "white", "line", True)
        d.text(x+18, y+14, 395, 27, label.upper(), 13, "orange", True)
        d.text(x+18, y+48, 395, 65, body, 18)
    d.text(38, 429, 884, 52, b[4] + "\n" + b[5], 15, "green", True)

    b = d.slide(8, "Verification / retained observations", "Tài")
    d.title(size=31)
    cards = [("2,324", "APPLICATION TESTS", "API Vitest 460 + Node 1,317\nWeb 547"),
             ("521/521", "CANONICAL PG17", "Managed-profile fixture 1/1\nreported separately"),
             ("45", "REAL HTTP CHECKS", "Docker + managed026\nGoogle login observed")]
    for i, (big, label, caption) in enumerate(cards):
        x = 38+i*300
        d.rect(x, 157, 284, 206, "white", "line", True)
        d.text(x+19, 177, 247, 67, big, 43, "green", family="Georgia")
        d.text(x+19, 254, 248, 22, label, 11, "orange", True)
        d.text(x+19, 292, 248, 55, caption, 16, "muted")
    d.text(40, 388, 875, 54, "Browser integration: real Nest/PG17, synthetic upstreams for selected workflows. Real-service evidence is identified separately.", 18)
    d.text(40, 455, 875, 29, b[4], 13, "muted", True)

    b = d.slide(9, "Evaluation / preserve the denominator", "Tài")
    d.title(size=31)
    d.text(40, 143, 872, 44, "Same model, corpus and prompt. Retrieval threshold: 0.55 → 0.50.", 18)
    labels = [("green", "PASS"), ("orange", "FAIL / provider error*"), ("sand", "INCONCLUSIVE"), ("gray", "NOT ATTEMPTED")]
    x = 222
    for c, label in labels:
        d.rect(x, 193, 12, 12, c)
        d.text(x+19, 190, 174, 20, label, 10, "muted")
        x += [100, 205, 173, 190][labels.index((c, label))]
    values = [[evidence['baseline'][k] for k in ['finalAnswerPass', 'finalAnswerFail', 'injectionInconclusive']] + [0],
              [evidence['improved'][k] for k in ['pass', 'blockedProvider503']] + [0, evidence['improved']['notAttempted']]]
    for i, vals in enumerate(values):
        y = 228+i*81
        d.text(38, y+7, 170, 35, "BASELINE / 20" if i == 0 else "FOLLOW-UP / 20", 13, bold=True)
        x = 222
        for val, (c, _) in zip(vals, labels):
            if not val:
                continue
            width = 700*val/20
            d.rect(x, y, width, 49, c)
            d.text(x, y+12, width, 30, str(val), 20, "white" if c in {"green", "orange"} else "ink", True, align="center")
            x += width
    d.text(222, 367, 697, 25, "*Baseline: 3 contract failures. Follow-up: 2 provider HTTP503 errors.", 11, "muted")
    d.rect(38, 404, 884, 78, "sand", radius=True)
    d.text(54, 416, 850, 31, b[4], 19, "ink", True)
    d.text(54, 451, 850, 26, "AI-assisted review. Injection: INCONCLUSIVE. Semantic-content 18/20 is a separate measure.", 12, "muted")

    b = d.slide(10, "Measurement / known versus unknown", "Tài")
    d.title(size=31)
    cards = [("4.58 s", "MEAN OPERATION", "20 baseline operations"),
             ("5.06 s", "SAMPLE p95", "Same 20; not a production SLA"),
             ("$0.0128272", "KNOWN GENERATION ESTIMATE", "20 generations with usage")]
    for i, (big, label, caption) in enumerate(cards):
        x = 38+i*300
        d.rect(x, 156, 284, 198, "white", "line", True)
        d.text(x+18, 179, 250, 60, big, 34 if i < 2 else 28, "green", family="Georgia")
        d.text(x+18, 254, 250, 21, label, 10, "orange", True)
        d.text(x+18, 294, 250, 50, caption, 15, "muted")
    d.text(38, 377, 884, 52, "20 of 40 baseline calls had missing usage. Unknown is not zero.\nCampaign upper used/reserved: USD 0.618895 across 51 calls, including a separate probe.", 17)
    d.text(38, 449, 884, 32, "Telemetry is best-effort. Estimates are not invoices or a provider hard spending cap.", 15, "green", True)

    b = d.slide(11, "Ownership / takeaways", "Tài + Thắng")
    d.title(size=32)
    for i in range(2):
        x = 38+i*451
        label, body = d.pair(b[i])
        d.rect(x, 149, 433, 181, "white", "line", True)
        d.text(x+20, 169, 390, 44, label, 29, "green", family="Georgia")
        d.text(x+20, 224, 391, 94, body, 19)
    d.text(38, 352, 880, 51, "Open: complete follow-up evaluation; broader semantic/device checks; personal confirmations.", 18)
    d.text(38, 423, 880, 51, "Thank you. Let’s trace a request together.", 28, "green", family="Georgia")

    b = d.slide(12, "Q&A backup / frontend", "Tài")
    d.title(size=32)
    for i in range(4):
        x = 38+i*224
        label, body = d.pair(b[i])
        d.rect(x, 155, 212, 213, "white", "line", True)
        d.text(x+14, 173, 183, 29, label, 18, "orange", True)
        d.text(x+14, 220, 183, 137, body, 18)
    d.text(38, 391, 880, 43, d.pair(b[4])[1], 18, "green", True)
    d.text(38, 449, 880, 34, d.pair(b[5])[1], 13, "muted")

    b = d.slide(13, "Q&A backup / code map", "Tài + Thắng")
    d.title(size=31)
    for i in range(4):
        label, body = d.pair(b[i])
        y = 150+i*68
        d.text(39, y+2, 176, 30, label, 16, "orange", True)
        d.text(221, y+2, 698, 47, body, 16)
        d.line(39, y+56, 921, y+56)
    d.text(39, 438, 882, 43, d.pair(b[5])[1], 18, "green", True)


def export_powerpoint(pptx, pdf):
    import win32com.client
    app = win32com.client.DispatchEx("PowerPoint.Application")
    presentation = None
    existing_count = app.Presentations.Count
    previous_security, previous_alerts = app.AutomationSecurity, app.DisplayAlerts
    try:
        app.AutomationSecurity = 3
        app.DisplayAlerts = 1
        presentation = app.Presentations.Open(str(pptx), ReadOnly=True, Untitled=False, WithWindow=False)
        overflow = []
        for number, slide in enumerate(presentation.Slides, 1):
            for shape in slide.Shapes:
                if shape.HasTextFrame and shape.TextFrame.HasText:
                    frame = shape.TextFrame2
                    text = frame.TextRange.Text
                    if frame.TextRange.BoundHeight > shape.Height + 2 or frame.TextRange.BoundWidth > shape.Width + 3:
                        overflow.append({'slide': number, 'text': text, 'boundHeight': frame.TextRange.BoundHeight,
                                         'boxHeight': shape.Height, 'boundWidth': frame.TextRange.BoundWidth,
                                         'boxWidth': shape.Width})
        if overflow:
            raise ValueError("PowerPoint text overflow: " + json.dumps(overflow, ensure_ascii=False))
        page_range = presentation.PrintOptions.Ranges.Add(1, presentation.Slides.Count)
        presentation.ExportAsFixedFormat(str(pdf), 2, PrintHiddenSlides=-1, PrintRange=page_range)
    finally:
        if presentation is not None:
            presentation.Close()
        app.AutomationSecurity, app.DisplayAlerts = previous_security, previous_alerts
        if existing_count == 0 and app.Presentations.Count == 0:
            app.Quit()


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    outline = load_outline()
    evidence = json.loads(EVIDENCE.read_text(encoding='utf8'))
    # Keep the frozen outcomes and sample denominators tied to the retained report.
    if [evidence['baseline'][k] for k in ['cases', 'finalAnswerPass', 'finalAnswerFail', 'injectionInconclusive']] != [20, 15, 3, 2]:
        raise ValueError('Evidence changed; revise content before exporting.')
    if [evidence['improved'][k] for k in ['attempted', 'pass', 'blockedProvider503', 'notAttempted']] != [3, 1, 2, 17]:
        raise ValueError('Follow-up evidence changed; review the narrative.')
    telemetry = evidence['observedTelemetry']['baseline']
    observed = [evidence['retainedApplicationTests']['total'],
                evidence['retainedDatabaseTests']['canonicalTests'],
                telemetry['operationMeanMs'], telemetry['operationP95Ms'],
                telemetry['knownGenerationCostUsd'], telemetry['completeUsageCalls'],
                telemetry['missingUsageCalls'], telemetry['providerCalls'],
                evidence['campaignIncludingSeparateLiteProbe']['upperUsedReservedUsd'],
                evidence['campaignIncludingSeparateLiteProbe']['providerCalls']]
    if observed != [2324, 521, 4583.45, 5056, .0128272, 20, 20, 40, '0.618895', 51]:
        raise ValueError('Retained metrics changed; review their labels and denominators.')
    deck = Deck(outline)
    compose(deck, evidence)
    for slide in list(deck.prs.slides)[12:]:
        slide._element.set('show', '0')
    bad = [g for g in deck.geometry if g['x'] < 0 or g['y'] < 0 or g['x']+g['w'] > W+.1 or g['y']+g['h'] > H+.1]
    if bad:
        raise ValueError('Geometry outside slide: ' + json.dumps(bad))
    pptx = out/(STEM+'.pptx')
    pdf = out/(STEM+'.pdf')
    deck.prs.save(pptx)
    export_powerpoint(pptx, pdf)
    with fitz.open(pdf) as doc:
        if len(doc) != 14:
            raise ValueError('Unexpected PDF page count')
        render = out/'preview'
        render.mkdir(exist_ok=True)
        for i, page in enumerate(doc):
            page.get_pixmap(matrix=fitz.Matrix(1.1, 1.1)).save(str(render/f'slide-{i+1:02}.png'))
        doc[0].get_pixmap(matrix=fitz.Matrix(1, 1)).save(str(out/(STEM+'-preview.png')))
    manifest = {
        'schema': 'examate-defense-deck-v1', 'date': '2026-10-24', 'slides': 14,
        'mainSlides': 12, 'backupSlides': [13, 14], 'aspectRatio': '16:9',
        'backupSlidesHiddenInPowerPoint': True,
        'onScreenLanguage': 'en', 'speakerNotesLanguage': 'vi', 'status': 'DRAFT',
        'statusReason': 'Incomplete improved evaluation and personal confirmations remain open.',
        'sourceBaseHead': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        'source': {'path': 'docs/SLIDES.md', 'sha256': sha(SOURCE)},
        'builder': {'path': 'scripts/build-defense-slides.py', 'sha256': sha(Path(__file__))},
        'retainedEvidence': {'path': 'docs/submission/report-evidence.json', 'sha256': sha(EVIDENCE)},
        'reportSource': {'path': 'docs/REPORT.md', 'sha256': sha(ROOT/'docs/REPORT.md')},
        'figureProvenance': {'path': 'docs/report-assets/figure-provenance.json', 'sha256': sha(ASSETS/'figure-provenance.json')},
        'assets': [{'path': str(p.relative_to(ROOT)).replace('\\', '/'), 'sha256': sha(p)} for p in sorted(deck.used_images)],
        'files': {p.name: {'bytes': p.stat().st_size, 'sha256': sha(p)} for p in [pptx, pdf, out/(STEM+'-preview.png')]},
        'editable': 'Text, diagrams and evaluation bars are native PowerPoint shapes; screenshots are retained raster assets.',
        'pdfMethod': 'Installed PowerPoint, hidden document, ExportAsFixedFormat including hidden backup slides; native text extents checked before export.',
        'noNewExecution': {'applicationTests': True, 'databaseTests': True, 'browser': True, 'provider': True, 'cloudWrites': True},
    }
    (out/'slides-manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False)+'\n', encoding='utf8')
    (out/'layout-audit.json').write_text(json.dumps(deck.geometry, indent=2, ensure_ascii=False), encoding='utf8')
    print(json.dumps({'slides': 14, 'powerPointOverflow': 0, 'pptx': str(pptx), 'pdf': str(pdf)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
