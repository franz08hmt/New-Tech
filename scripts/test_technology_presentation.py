"""Inspect exported PDF and PPTX artifacts rather than source alone."""
import hashlib
import json
from pathlib import Path
import unittest

import fitz
from bs4 import BeautifulSoup
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/submission"
STEM = "ExaMate-Technologies-VI-20261026"


class TechnologyArtifactTests(unittest.TestCase):
    def test_pdf_preserves_all_technology_topics_and_vietnamese(self):
        required = ["Công nghệ", "Guard", "TypeScript", "RLS", "PDF",
                    "JSON", "Whisper", "admission", "Đóng gói", "bằng chứng"]
        with fitz.open(OUTPUT / f"{STEM}.pdf") as document:
            self.assertEqual(len(document), 10)
            for page, text in zip(document, required):
                with self.subTest(page=page.number + 1):
                    self.assertAlmostEqual(page.rect.width / page.rect.height, 16 / 9, places=3)
                    self.assertIn(text.casefold(), page.get_text().casefold())
                    self.assertNotIn("\ufffd", page.get_text())
            self.assertTrue(any("Poppins" in f[3] for f in document[0].get_fonts()))
            self.assertIn("17 case chưa chạy", document[-1].get_text().replace("\n", " "))

    def test_pptx_has_notes_and_identifies_layout_preserving_artwork(self):
        presentation = Presentation(OUTPUT / f"{STEM}.pptx")
        self.assertEqual(len(presentation.slides), 10)
        self.assertAlmostEqual(presentation.slide_width / presentation.slide_height, 16 / 9, places=5)
        self.assertEqual(presentation.core_properties.identifier,
                         hashlib.sha256((OUTPUT / f"{STEM}.pdf").read_bytes()).hexdigest())
        for slide in presentation.slides:
            self.assertEqual(len(slide.shapes), 1)
            self.assertIn("Bố cục cố định", slide.notes_slide.notes_text_frame.text)
            self.assertEqual(len(slide._element.xpath("./p:transition/p:fade")), 1)

    def test_manifest_and_original_screenshot_binding(self):
        manifest = json.loads((OUTPUT / "technology-deck-manifest.json").read_text(encoding="utf-8"))
        for key in ["html", "pdf", "pptx", "preview", "knowledge"]:
            self.assertEqual(manifest[f"{key}Sha256"],
                             hashlib.sha256((OUTPUT / manifest[key]).read_bytes()).hexdigest())
        source = BeautifulSoup((ROOT / "docs/presentation/technology.html").read_text(encoding="utf-8"), "html.parser")
        self.assertEqual(len(source.select(".speaker")), 10)

    def test_knowledge_companion_preserves_all_original_details(self):
        data = json.loads((ROOT / "docs/presentation/technology-knowledge.json").read_text(encoding="utf-8"))
        companion = BeautifulSoup((OUTPUT / "ExaMate-Technologies-Knowledge-VI-20261026.html").read_text(encoding="utf-8"), "html.parser")
        normalize = lambda value: " ".join(value.split())
        self.assertEqual(data["sourceCheckpoint"], "69eab36")
        self.assertEqual(len(data["sections"]), 15)
        for section in data["sections"]:
            rendered = normalize(companion.find(id=section["id"]).get_text(" ", strip=True))
            for detail in [*section["content"], section["notes"]]:
                self.assertIn(normalize(detail), rendered)


if __name__ == "__main__":
    unittest.main()
